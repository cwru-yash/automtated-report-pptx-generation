import logging
import os
from pathlib import Path
from enum import Enum
from typing import Dict, List, Literal

from pydantic import BaseModel

logger = logging.getLogger(__name__)


class PlaceholderNotFoundError(Exception):
    """Raised when a required placeholder is missing from a template file."""
    pass


class TemplateType(str, Enum):
    ppt = "ppt"
    html = "html"


class PlaceholderSpec(BaseModel):
    name: str
    type: Literal["text", "image", "number"]


class TemplateSpec(BaseModel):
    template_id: str
    file_path: str
    template_type: TemplateType
    placeholders: List[PlaceholderSpec] = []


class TemplateRegistry:
    def __init__(self):
        self._registry: Dict[str, TemplateSpec] = {}

    def register(self, spec: TemplateSpec) -> None:
        if spec.template_id in self._registry:
            raise ValueError(f"Template already registered: {spec.template_id}")
        self._registry[spec.template_id] = spec
        logger.info(f"Registered template: {spec.template_id} ({spec.template_type})")

    def clear(self) -> None:
        self._registry.clear()

    def get_template(self, template_id: str) -> TemplateSpec:
        if template_id not in self._registry:
            raise KeyError(f"Template not found: {template_id}")
        return self._registry[template_id]

    def list_templates(self) -> List[TemplateSpec]:
        return list(self._registry.values())

    def validate_ppt_placeholders(self, spec: TemplateSpec) -> None:
        """Validates that all required placeholders exist as shape names in the .pptx file."""
        if not spec.placeholders:
            return
        try:
            from pptx import Presentation
        except ImportError:
            logger.warning("python-pptx not installed — skipping placeholder validation")
            return

        prs = Presentation(spec.file_path)
        shape_names = set()
        for slide in prs.slides:
            for shape in slide.shapes:
                shape_names.add(shape.name)

        for ph in spec.placeholders:
            if ph.name not in shape_names:
                raise PlaceholderNotFoundError(
                    f"Placeholder '{ph.name}' not found in template '{spec.template_id}'. "
                    f"Available shapes: {sorted(shape_names)}"
                )


# Module-level singleton
registry = TemplateRegistry()


def _template_id_for(path: Path, root: Path, existing_ids: set[str]) -> str:
    template_id = path.stem
    if template_id not in existing_ids:
        return template_id
    relative = path.relative_to(root).with_suffix("")
    return "__".join(relative.parts)


def init_registry(templates_dir: str) -> None:
    """Scans templates_dir recursively for .pptx files and registers each as a TemplateSpec."""
    root = Path(templates_dir)
    if not root.is_dir():
        logger.warning(f"Templates directory not found: {templates_dir} — skipping registry init")
        return

    count = 0
    seen_ids = {template.template_id for template in registry.list_templates()}
    for path in sorted(root.rglob("*.pptx")):
        if path.name.startswith("~$"):
            continue

        template_id = _template_id_for(path, root, seen_ids)
        schema_path = path.with_name(f"{path.stem}.schema.json")
        placeholders: List[PlaceholderSpec] = []
        if schema_path.is_file():
            import json
            with open(schema_path, encoding="utf-8") as f:
                raw = json.load(f)
            placeholders = [PlaceholderSpec(**ph) for ph in raw.get("placeholders", [])]

        spec = TemplateSpec(
            template_id=template_id,
            file_path=str(path),
            template_type=TemplateType.ppt,
            placeholders=placeholders,
        )
        try:
            registry.register(spec)
            seen_ids.add(template_id)
            count += 1
        except ValueError:
            logger.debug(f"Template already registered (skipped): {template_id}")

    logger.info(f"Template registry initialized: {count} template(s) loaded from '{templates_dir}'")
