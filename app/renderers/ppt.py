"""
PPT Renderer — Professional report generation using the client_cvc_master template.

Strategy:
1. Open the client_cvc_master.pptx (58 slides, branded).
2. Inject LLM narrative text into slides 2–6 (named placeholders:
   exec_summary, methodology, brand_health, weak_links, takeaways).
3. Append one chart slide per available analysis chart, using the
   template's built-in Blank Page: White layout for consistency.
"""
from __future__ import annotations

import logging
import os
from typing import Any, Dict, Optional

from app.bundle.schema import AnalysisBundle
from app.deck.layout_schema import ContentBlock, DeckDocument, DeckSlideDocument
from app.templates.registry import TemplateType, registry

logger = logging.getLogger(__name__)

# Map section name → (slide_index_0based, shape_name) in client_cvc_master
_SECTION_SLIDE_MAP: dict[str, tuple[int, str]] = {
    "exec_summary": (1, "exec_summary"),
    "methodology":  (2, "methodology"),
    "brand_health": (3, "brand_health"),
    "weak_links":   (4, "weak_links"),
    "takeaways":    (5, "takeaways"),
}

_CHART_TITLES: dict[int, str] = {
    1: "Industry Average Curve",
    2: "Strong Links — Industry",
    3: "Weak Links — Industry",
    4: "S&R Index Ranking",
    5: "Strong Links — Brand",
    6: "Achilles Heel — Brand",
    7: "Best of Best",
    8: "Breakdown by Segment",
}


class PPTRenderer:
    def render(
        self,
        bundle: AnalysisBundle,
        sections: Dict[str, str],
        template_id: str,
        output_path: str,
        images: Optional[Dict[str, str]] = None,
        deck_plan: Optional[Any] = None,
    ) -> str:
        """
        Renders a .pptx file by:
        1. Injecting narrative sections into named shape placeholders on slides 2–6.
        2. Appending analysis chart slides after the main content.
        """
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt, Emu
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN
        except ImportError:
            raise ImportError("python-pptx is required. Install with: poetry add python-pptx")

        images = images or {}
        spec = registry.get_template(template_id)
        if spec.template_type != TemplateType.ppt:
            raise ValueError(f"Template '{template_id}' is not a PPT template")

        prs = Presentation(spec.file_path)

        # --- Step 1: Inject text into named placeholder shapes (slides 2–6) ---
        for section_name, content in sections.items():
            if section_name not in _SECTION_SLIDE_MAP:
                continue
            slide_idx, shape_name = _SECTION_SLIDE_MAP[section_name]
            if slide_idx >= len(prs.slides):
                logger.warning("Slide index %d out of range for section '%s'", slide_idx, section_name)
                continue
            slide = prs.slides[slide_idx]
            injected = False
            for shape in slide.shapes:
                if shape.name == shape_name and getattr(shape, "has_text_frame", False):
                    self._inject_text(shape, content)
                    logger.info("Injected '%s' → slide %d shape '%s'", section_name, slide_idx + 1, shape_name)
                    injected = True
                    break
            if not injected:
                logger.warning("Shape '%s' not found on slide %d for section '%s'", shape_name, slide_idx + 1, section_name)

        if deck_plan:
            self._append_deck_plan_slides(prs, deck_plan, images)
            prs.save(output_path)
            logger.info("PPT rendered to %s", output_path)
            return output_path

        # --- Step 2: Append one chart slide per analysis chart ---
        blank_layout = self._find_blank_layout(prs)

        # Sort chart keys: "chart_1", "chart_2", ... or plain int keys
        sorted_chart_keys = sorted(
            images.keys(),
            key=lambda k: int(k.replace("chart_", "")) if k.replace("chart_", "").isdigit() else 999
        )

        for chart_key in sorted_chart_keys:
            image_path = images[chart_key]
            if not os.path.isfile(image_path):
                logger.warning("Chart image not found: %s", image_path)
                continue

            # Resolve analysis type int for title lookup
            atype_str = chart_key.replace("chart_", "")
            atype = int(atype_str) if atype_str.isdigit() else 0
            title = _CHART_TITLES.get(atype, chart_key.replace("_", " ").title())

            slide = prs.slides.add_slide(blank_layout)
            self._add_chart_slide_content(slide, title, image_path)
            logger.info("Appended chart slide: type=%d title='%s'", atype, title)

        prs.save(output_path)
        logger.info("PPT rendered to %s", output_path)
        return output_path

    def render_deck_document(
        self,
        deck_document: DeckDocument | dict[str, Any],
        template_id: str,
        output_path: str,
        images: Optional[Dict[str, str]] = None,
    ) -> str:
        """Render validated editable DeckDocument JSON to a semantic PPT fallback."""
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError("python-pptx is required. Install with: poetry add python-pptx")

        images = images or {}
        deck = DeckDocument.model_validate(deck_document)
        spec = registry.get_template(template_id)
        if spec.template_type != TemplateType.ppt:
            raise ValueError(f"Template '{template_id}' is not a PPT template")

        prs = Presentation(spec.file_path)
        blank_layout = self._find_blank_layout(prs)
        for deck_slide in deck.slides:
            slide = prs.slides.add_slide(blank_layout)
            self._add_deck_document_slide_content(slide, deck_slide, images)

        prs.save(output_path)
        logger.info("DeckDocument PPT rendered to %s", output_path)
        return output_path

    def _inject_text(self, shape, content: str) -> None:
        """Replace all text in a shape's text frame, preserving paragraph formatting."""
        from pptx.util import Pt

        tf = shape.text_frame
        tf.word_wrap = True

        # Clear existing content
        for para in tf.paragraphs:
            for run in para.runs:
                run.text = ""

        # Use the first paragraph to inject text
        if tf.paragraphs:
            para = tf.paragraphs[0]
            if para.runs:
                para.runs[0].text = content
            else:
                run = para.add_run()
                run.text = content
        else:
            # Fallback: clear and set directly
            tf.clear()
            tf.paragraphs[0].text = content

    def _find_blank_layout(self, prs):
        """Find the best blank or minimal slide layout for chart slides."""
        preferred = ["Blank Page: White", "Blank", "blank"]
        for layout in prs.slide_layouts:
            if layout.name in preferred:
                return layout
        # Fallback: layout with fewest placeholders
        return min(prs.slide_layouts, key=lambda l: len(l.placeholders))

    def _append_deck_plan_slides(self, prs, deck_plan: Any, images: Dict[str, str]) -> None:
        """Append generated consulting slides from a structured DeckPlan."""
        if hasattr(deck_plan, "model_dump"):
            deck_plan = deck_plan.model_dump()

        blank_layout = self._find_blank_layout(prs)
        for planned_slide in deck_plan.get("slides", []):
            slide = prs.slides.add_slide(blank_layout)
            self._add_planned_slide_content(slide, planned_slide, images)

    def _add_planned_slide_content(
        self,
        slide,
        planned_slide: Dict[str, Any],
        images: Dict[str, str],
    ) -> None:
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        title = str(planned_slide.get("title") or "Untitled Slide")
        subtitle = str(planned_slide.get("subtitle") or "")
        bullets = [str(item) for item in planned_slide.get("bullets") or []]
        evidence = [str(item) for item in planned_slide.get("evidence") or []]
        chart_refs = [str(item) for item in planned_slide.get("chart_refs") or []]
        chart_path = next(
            (
                images.get(ref)
                for ref in chart_refs
                if images.get(ref) and os.path.isfile(images[ref])
            ),
            None,
        )

        title_box = slide.shapes.add_textbox(
            Inches(0.55), Inches(0.25), Inches(12.2), Inches(0.55)
        )
        title_frame = title_box.text_frame
        title_frame.clear()
        title_run = title_frame.paragraphs[0].add_run()
        title_run.text = title
        title_run.font.size = Pt(23)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

        if subtitle:
            subtitle_box = slide.shapes.add_textbox(
                Inches(0.58), Inches(0.82), Inches(12.0), Inches(0.35)
            )
            subtitle_frame = subtitle_box.text_frame
            subtitle_frame.clear()
            subtitle_run = subtitle_frame.paragraphs[0].add_run()
            subtitle_run.text = subtitle
            subtitle_run.font.size = Pt(12)
            subtitle_run.font.color.rgb = RGBColor(0x5B, 0x65, 0x72)

        if chart_path:
            slide.shapes.add_picture(
                chart_path,
                Inches(0.55),
                Inches(1.35),
                width=Inches(7.6),
                height=Inches(4.5),
            )
            bullet_box = slide.shapes.add_textbox(
                Inches(8.45), Inches(1.35), Inches(4.05), Inches(4.5)
            )
        else:
            bullet_box = slide.shapes.add_textbox(
                Inches(0.85), Inches(1.35), Inches(11.5), Inches(4.45)
            )

        self._write_bullets(bullet_box.text_frame, bullets)

        if evidence:
            evidence_box = slide.shapes.add_textbox(
                Inches(0.85), Inches(6.05), Inches(11.5), Inches(0.8)
            )
            evidence_frame = evidence_box.text_frame
            evidence_frame.clear()
            heading = evidence_frame.paragraphs[0].add_run()
            heading.text = "Evidence: "
            heading.font.size = Pt(10)
            heading.font.bold = True
            heading.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

            evidence_run = evidence_frame.paragraphs[0].add_run()
            evidence_run.text = " ".join(evidence[:3])
            evidence_run.font.size = Pt(10)
            evidence_run.font.color.rgb = RGBColor(0x5B, 0x65, 0x72)

    def _write_bullets(self, text_frame, bullets: list[str]) -> None:
        from pptx.util import Pt
        from pptx.dml.color import RGBColor

        text_frame.clear()
        if not bullets:
            bullets = ["No slide bullets were generated."]

        for index, bullet in enumerate(bullets[:5]):
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(14)
            paragraph.font.color.rgb = RGBColor(0x2F, 0x3B, 0x4A)
            paragraph.space_after = Pt(8)

    def _add_chart_slide_content(self, slide, title: str, image_path: str) -> None:
        """Add a title textbox and full-bleed chart image to a chart slide."""
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        # Title bar at top
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.2), Inches(12.3), Inches(0.6)
        )
        tf = title_box.text_frame
        tf.clear()
        para = tf.paragraphs[0]
        run = para.add_run()
        run.text = title
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

        # Chart image — large, centered below title
        slide.shapes.add_picture(
            image_path,
            Inches(0.5),
            Inches(0.95),
            width=Inches(12.3),
            height=Inches(6.2),
        )

    def _add_deck_document_slide_content(
        self,
        slide,
        deck_slide: DeckSlideDocument,
        images: Dict[str, str],
    ) -> None:
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        title_box = slide.shapes.add_textbox(
            Inches(0.55), Inches(0.25), Inches(12.2), Inches(0.55)
        )
        title_frame = title_box.text_frame
        title_frame.clear()
        title_run = title_frame.paragraphs[0].add_run()
        title_run.text = deck_slide.title
        title_run.font.size = Pt(23)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

        body_box = slide.shapes.add_textbox(
            Inches(0.75), Inches(1.05), Inches(11.8), Inches(5.75)
        )
        body_frame = body_box.text_frame
        body_frame.word_wrap = True
        body_frame.clear()

        lines = self._deck_block_lines(deck_slide.content, images)
        if not lines:
            lines = ["No editable deck content was available for this slide."]

        for index, line in enumerate(lines[:18]):
            paragraph = body_frame.paragraphs[0] if index == 0 else body_frame.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(12 if line.startswith(("•", "  ")) else 13)
            paragraph.font.color.rgb = RGBColor(0x2F, 0x3B, 0x4A)
            paragraph.space_after = Pt(5)

    def _deck_block_lines(
        self,
        blocks: list[ContentBlock],
        images: Dict[str, str],
    ) -> list[str]:
        lines: list[str] = []
        for block in blocks:
            lines.extend(self._single_deck_block_lines(block, images))
        return lines

    def _single_deck_block_lines(
        self,
        block: ContentBlock,
        images: Dict[str, str],
    ) -> list[str]:
        if block.type == "heading":
            return [str(block.text)]
        if block.type in {"paragraph", "quote", "callout", "footnote", "source_note"}:
            prefix = {
                "quote": "Quote: ",
                "callout": "Callout: ",
                "footnote": "Footnote: ",
                "source_note": "",
            }.get(block.type, "")
            return [f"{prefix}{block.text}"]
        if block.type in {"bullet_list", "ordered_list"}:
            return [f"{index + 1}. {item}" if block.type == "ordered_list" else f"• {item}" for index, item in enumerate(block.items)]
        if block.type == "metric_card":
            metric = f"{block.label}: {block.value}"
            return [metric, f"  {block.helper}"] if block.helper else [metric]
        if block.type == "table":
            header_line = "Table: " + " | ".join(str(header) for header in block.headers)
            row_lines = ["  " + " | ".join(str(cell) for cell in row) for row in block.rows[:4]]
            return [header_line, *row_lines]
        if block.type == "chart_placeholder":
            chart_title = f" — {block.title}" if block.title else ""
            chart_status = "available" if images.get(block.chart_ref or "") else "placeholder"
            return [f"Chart {chart_status}: {block.chart_ref}{chart_title}"]
        if block.type == "image":
            return [f"Image placeholder: {block.alt or block.src}"]
        if block.type in {"two_column", "three_column"}:
            lines: list[str] = []
            for index, column in enumerate(block.columns, start=1):
                lines.append(f"Column {index}")
                lines.extend(self._deck_block_lines(column.content, images))
            return lines
        if block.type == "divider":
            return ["---"]
        return [f"Unsupported block type: {block.type}"]


def render_ppt(
    bundle: AnalysisBundle,
    sections: Dict[str, str],
    template_id: str,
    output_path: str,
    images: Optional[Dict[str, str]] = None,
    deck_plan: Optional[Any] = None,
) -> str:
    return PPTRenderer().render(bundle, sections, template_id, output_path, images, deck_plan)


def render_deck_document_ppt(
    deck_document: DeckDocument | dict[str, Any],
    template_id: str,
    output_path: str,
    images: Optional[Dict[str, str]] = None,
) -> str:
    return PPTRenderer().render_deck_document(deck_document, template_id, output_path, images)
