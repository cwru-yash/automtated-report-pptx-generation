import asyncio
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any, Dict, Literal, Optional
from uuid import uuid4

from fastapi import FastAPI
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
import logging
from pydantic import BaseModel, Field

from app.config import settings
from app.core.database import init_db
from app.deck.routes import router as deck_router
from app.observability import StructuredEventLogger, read_jsonl
from app.templates.registry import init_registry, registry

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class ReportJobRequest(BaseModel):
    wave_id: str = Field(..., description="Completed Decomposer wave ID")
    language: str = Field(default="en-US", description="Locale for report labels and narrative")
    template_id: Optional[str] = Field(
        default="client_cvc_master",
        description="Registered PPT template ID. Use null to skip PPT rendering.",
    )
    deck_mode: Literal["planned", "legacy"] = Field(
        default="planned",
        description="Use deterministic evidence-backed deck planning or legacy narrative rendering",
    )
    allow_partial: bool = Field(
        default=False,
        description="Allow diagnostic output from incomplete wave data",
    )
    audience: str = Field(default="executive stakeholders", description="Target deck audience")
    tone: str = Field(default="consulting", description="Deck communication tone")
    use_llm: bool = Field(default=False, description="Use LLM-backed LangGraph narrative nodes in legacy mode")
    require_llm: bool = Field(default=False, description="Fail sections instead of using fallback text")
    include_pdf: bool = Field(default=True, description="Render PDF artifact when WeasyPrint is available")
    output_dir: str = Field(default="artifacts", description="Local output directory for demo artifacts")


class ReportJobResponse(BaseModel):
    job_id: str
    batch_id: Optional[str] = None
    status: str
    bundle: Dict[str, Any]
    sections: Dict[str, str]
    ppt_sections: Dict[str, str] = Field(default_factory=dict)
    llm_metadata: Dict[str, Any] = Field(default_factory=dict)
    artifacts: Dict[str, str]
    warnings: list[str] = Field(default_factory=list)
    logs_url: str


class BatchReportJobRequest(BaseModel):
    wave_id: str = Field(..., description="Completed Decomposer wave ID")
    languages: list[str] = Field(
        default_factory=lambda: ["en-US"],
        description="Locales to generate as independent child jobs.",
    )
    template_id: Optional[str] = Field(
        default="client_cvc_master",
        description="Registered PPT template ID. Use null to skip PPT rendering.",
    )
    deck_mode: Literal["planned", "legacy"] = Field(
        default="planned",
        description="Use deterministic evidence-backed deck planning or legacy narrative rendering",
    )
    allow_partial: bool = Field(
        default=False,
        description="Allow diagnostic output from incomplete wave data",
    )
    audience: str = Field(default="executive stakeholders", description="Target deck audience")
    tone: str = Field(default="consulting", description="Deck communication tone")
    use_llm: bool = Field(default=False, description="Use LLM-backed LangGraph narrative nodes in legacy mode")
    require_llm: bool = Field(default=False, description="Fail sections instead of using fallback text")
    include_pdf: bool = Field(default=True, description="Render PDF artifact when WeasyPrint is available")
    output_dir: str = Field(default="artifacts", description="Local output directory for demo artifacts")


class BatchChildJobResponse(BaseModel):
    language: str
    job_id: str
    status: str
    artifacts: Dict[str, str] = Field(default_factory=dict)
    warnings: list[str] = Field(default_factory=list)
    llm_metadata: Dict[str, Any] = Field(default_factory=dict)
    logs_url: str


class BatchReportJobResponse(BaseModel):
    batch_id: str
    status: str
    wave_id: str
    template_id: Optional[str]
    languages: list[str]
    children: list[BatchChildJobResponse]
    logs_url: str


def ensure_template_registry_loaded() -> None:
    if not registry.list_templates():
        init_registry(settings.TEMPLATES_DIR)


def dedupe_preserving_order(values: list[str]) -> list[str]:
    return list(dict.fromkeys(values))


def log_path(kind: str, identifier: str) -> Path:
    return Path(settings.LOGS_DIR) / kind / f"{identifier}.jsonl"


def get_demo_dependencies():
    BundleAssembler = getattr(app.state, "bundle_assembler_cls", None)
    run_narrative_engine_async = getattr(app.state, "run_narrative_engine_async", None)
    render_ppt = getattr(app.state, "render_ppt", None)
    html_renderer = getattr(app.state, "html_renderer", None)
    chart_generator = getattr(app.state, "chart_generator", None)
    wave_provider = getattr(app.state, "wave_provider", None)

    if not all([BundleAssembler, run_narrative_engine_async, render_ppt, html_renderer]):
        from app.bundle.assembler import BundleAssembler
        from app.narrative.graph import run_narrative_engine_async
        from app.renderers.html import HTMLRenderer
        from app.renderers.ppt import render_ppt

        html_renderer = HTMLRenderer(settings.TEMPLATES_DIR, settings.LOCALES_DIR)

    if chart_generator is None:
        from app.charts.generator import ChartGenerator

        chart_generator = ChartGenerator()
        app.state.chart_generator = chart_generator

    if wave_provider is None:
        from app.data.provider import get_provider

        wave_provider = get_provider(settings)
        app.state.wave_provider = wave_provider

    return BundleAssembler, run_narrative_engine_async, render_ppt, html_renderer, chart_generator, wave_provider


def prewarm_demo_dependencies(app: FastAPI) -> None:
    """
    Load the demo generation path during startup so the first recorded job call
    does not pay import and template-cache costs on camera.
    """
    try:
        from app.bundle.assembler import BundleAssembler
        from app.bundle.schema import AnalysisBundle
        from app.narrative.graph import run_narrative_engine_async
        from app.renderers.html import HTMLRenderer
        from app.renderers.ppt import render_ppt
        from app.charts.generator import ChartGenerator
        from app.data.provider import get_provider

        app.state.bundle_assembler_cls = BundleAssembler
        app.state.run_narrative_engine_async = run_narrative_engine_async
        app.state.render_ppt = render_ppt
        app.state.html_renderer = HTMLRenderer(settings.TEMPLATES_DIR, settings.LOCALES_DIR)
        app.state.chart_generator = ChartGenerator()
        app.state.wave_provider = get_provider(settings)

        bundle = AnalysisBundle(
            wave_id="startup-prewarm",
            language="en-US",
            metrics={"nps": 0.0, "brand_awareness": 0.0, "cvc_score": 0.0},
            analysis_results={},
        )
        sections = {
            "exec_summary": "Prewarm",
            "methodology": "Prewarm",
            "brand_health": "Prewarm",
            "weak_links": "Prewarm",
            "takeaways": "Prewarm",
        }
        app.state.html_renderer.render(bundle, sections, language="en-US")

        try:
            from pptx import Presentation

            template = registry.get_template("demo_master")
            Presentation(template.file_path)
        except Exception as exc:
            logger.warning("PPT prewarm skipped: %s", exc)
    except Exception as exc:
        logger.warning("Demo dependency prewarm skipped: %s", exc)

@asynccontextmanager
async def lifespan(app: FastAPI):
    logger.info("Application starting up... Settings loaded securely.")
    init_registry(settings.TEMPLATES_DIR)
    await init_db()
    prewarm_demo_dependencies(app)
    yield
    logger.info("Application shutting down...")

app = FastAPI(title="Automated Report Platform", lifespan=lifespan)
STATIC_DIR = Path(__file__).parent / "static"
ARTIFACTS_DIR = Path("artifacts")
DECK_EDITOR_DIST_DIR = Path("frontend") / "dist"
ARTIFACTS_DIR.mkdir(exist_ok=True)

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/artifacts", StaticFiles(directory=ARTIFACTS_DIR), name="artifacts")
app.mount(
    "/decks/editor/assets",
    StaticFiles(directory=DECK_EDITOR_DIST_DIR / "assets", check_dir=False),
    name="deck-editor-assets",
)

app.include_router(deck_router)


@app.get("/", include_in_schema=False)
def report_console():
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/decks/editor", include_in_schema=False)
@app.get("/decks/editor/{deck_id}", include_in_schema=False)
def deck_editor(deck_id: str | None = None):
    index_path = DECK_EDITOR_DIST_DIR / "index.html"
    if index_path.is_file():
        return FileResponse(index_path)
    return FileResponse(STATIC_DIR / "index.html")

@app.get("/health")
def health_check():
    return {"status": "ok", "service": "automated-report-platform"}


@app.get("/api/v1/locales")
def list_locales():
    locales_dir = Path(settings.LOCALES_DIR)
    locales = []
    if locales_dir.is_dir():
        locales = sorted(path.stem for path in locales_dir.glob("*.json"))
    return {"locales": locales}


@app.get("/api/v1/templates")
def list_templates():
    ensure_template_registry_loaded()
    return {"templates": [t.model_dump() for t in registry.list_templates()]}


@app.get("/api/v1/templates/{template_id}/inspect")
def inspect_template(template_id: str):
    ensure_template_registry_loaded()
    spec = registry.get_template(template_id)
    if spec.template_type.value != "ppt":
        return {"template": spec.model_dump(), "slides": []}

    from pptx import Presentation

    prs = Presentation(spec.file_path)
    required = {placeholder.name for placeholder in spec.placeholders}
    seen = set()
    slides = []

    for slide_number, slide in enumerate(prs.slides, start=1):
        shapes = []
        for shape in slide.shapes:
            text = ""
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip().replace("\n", " ")
            seen.add(shape.name)
            shapes.append(
                {
                    "name": shape.name,
                    "type": str(shape.shape_type),
                    "has_text_frame": bool(getattr(shape, "has_text_frame", False)),
                    "text_sample": text[:100],
                }
            )
        slides.append({"slide_number": slide_number, "shapes": shapes})

    return {
        "template": spec.model_dump(),
        "required_placeholders": sorted(required),
        "missing_placeholders": sorted(required - seen),
        "slides": slides,
    }


async def execute_report_job(
    request: ReportJobRequest,
    *,
    job_id: str | None = None,
    batch_id: str | None = None,
) -> ReportJobResponse:
    ensure_template_registry_loaded()
    BundleAssembler, run_narrative_engine_async, render_ppt, html_renderer, chart_generator, wave_provider = get_demo_dependencies()

    job_id = job_id or str(uuid4())
    output_dir = Path(request.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    event_log = StructuredEventLogger(
        settings.LOGS_DIR,
        batch_id=batch_id,
        job_id=job_id,
        context={
            "wave_id": request.wave_id,
            "language": request.language,
            "template_id": request.template_id,
            "deck_mode": request.deck_mode,
        },
    )
    event_log.emit(
        "job.started",
        deck_mode=request.deck_mode,
        use_llm=request.use_llm,
        require_llm=request.require_llm,
        include_pdf=request.include_pdf,
    )

    try:
        wave_data = wave_provider.get_wave_data(request.wave_id)
        event_log.emit(
            "wave_data.loaded",
            wave_id_in_payload=wave_data.get("wave", {}).get("id", ""),
            analysis_count=len(wave_data.get("analyses", [])),
        )
        bundle = await BundleAssembler().build_bundle(
            request.wave_id,
            request.language,
            wave_data=wave_data,
        )
        
        # Inject provenance metadata
        metadata = wave_data.get("metadata", {})
        bundle.gold_activity_version_id = metadata.get("gold_activity_version_id", "")
        bundle.analysis_run_id = metadata.get("analysis_run_id", "")

        event_log.emit(
            "bundle.assembled",
            metric_keys=sorted(bundle.metrics.keys()),
            analysis_keys=sorted(bundle.analysis_results.keys()),
            provenance={
                "gold_activity_version_id": bundle.gold_activity_version_id,
                "analysis_run_id": bundle.analysis_run_id
            }
        )

        deck_plan = None
        warnings: list[str] = []

        if request.deck_mode == "planned":
            from app.deck.extractor import FindingExtractor
            from app.deck.planner import SlidePlanner
            from app.deck.schema import DeckDataQualityError, DeckValidationError

            try:
                findings_context = FindingExtractor().extract(
                    bundle,
                    wave_data,
                    allow_partial=request.allow_partial,
                )
                deck_plan = SlidePlanner().create_plan(
                    findings_context,
                    bundle,
                    audience=request.audience,
                    tone=request.tone,
                )
                sections = deck_plan.as_sections()
                ppt_sections = sections
                llm_metadata = {
                    "deck_mode": "planned",
                    "llm_used": False,
                    "finding_count": len(findings_context.findings),
                    "slide_count": len(deck_plan.slides),
                }
                warnings = list(deck_plan.validation_warnings)
                if request.use_llm:
                    warnings.append(
                        "LLM request ignored for planned deck mode; deterministic planner used."
                    )
                status = "completed"
                event_log.emit(
                    "deck_plan.completed",
                    status=status,
                    finding_count=len(findings_context.findings),
                    slide_count=len(deck_plan.slides),
                    warning_count=len(warnings),
                )
            except (DeckDataQualityError, DeckValidationError) as exc:
                status = "failed"
                warnings = [str(exc)]
                event_log.emit(
                    "job.failed",
                    level="warning",
                    reason="deck_plan_blocked",
                    warning_count=len(warnings),
                    error=str(exc),
                )
                return ReportJobResponse(
                    job_id=job_id,
                    batch_id=batch_id,
                    status=status,
                    bundle=bundle.model_dump(),
                    sections={},
                    ppt_sections={},
                    llm_metadata={"deck_mode": "planned", "llm_used": False},
                    artifacts={},
                    warnings=warnings,
                    logs_url=f"/api/v1/jobs/{job_id}/logs",
                )
        else:
            narrative = await run_narrative_engine_async(
                bundle,
                use_llm=request.use_llm,
                require_llm=request.require_llm,
            )
            sections = narrative.get("sections", {})
            ppt_sections = narrative.get("ppt_sections", {})
            llm_metadata = narrative.get("llm_metadata", {})
            warnings = [str(error) for error in narrative.get("errors", [])]
            status = "failed" if request.require_llm and warnings else "completed"
            event_log.emit(
                "narrative.completed",
                status=status,
                section_count=len(sections),
                llm_modes={
                    key: value.get("mode")
                    for key, value in llm_metadata.items()
                },
                warning_count=len(warnings),
            )

        artifacts: Dict[str, str] = {}

        if status == "failed":
            event_log.emit(
                "job.failed",
                level="warning",
                reason="required_llm_warning",
                warning_count=len(warnings),
            )
            return ReportJobResponse(
                job_id=job_id,
                batch_id=batch_id,
                status=status,
                bundle=bundle.model_dump(),
                sections=sections,
                ppt_sections=ppt_sections,
                llm_metadata=llm_metadata,
                artifacts=artifacts,
                warnings=warnings,
                logs_url=f"/api/v1/jobs/{job_id}/logs",
            )

        if deck_plan:
            slide_plan_path = output_dir / f"{job_id}_slide_plan.json"
            slide_plan_path.write_text(deck_plan.model_dump_json(indent=2), encoding="utf-8")
            artifacts["slide_plan"] = str(slide_plan_path)
            event_log.emit(
                "artifact.written",
                artifact_type="slide_plan",
                path=str(slide_plan_path),
            )

        html = html_renderer.render(bundle, sections, language=request.language)
        html_path = output_dir / f"{job_id}.html"
        html_path.write_text(html, encoding="utf-8")
        artifacts["html"] = str(html_path)
        event_log.emit("artifact.written", artifact_type="html", path=str(html_path))

        # Generate charts
        charts = chart_generator.render_all(
            wave_data,
            skip_empty=request.deck_mode == "planned",
        )
        chart_paths = {}
        for analysis_type, png_bytes in charts.items():
            cpath = output_dir / f"{job_id}_chart_{analysis_type}.png"
            cpath.write_bytes(png_bytes)
            chart_paths[f"chart_{analysis_type}"] = str(cpath)
        
        if chart_paths:
            event_log.emit("charts.rendered", count=len(chart_paths))

        if deck_plan:
            expected_charts = {
                chart_ref
                for slide in deck_plan.slides
                for chart_ref in slide.chart_refs
            }
            missing_charts = sorted(expected_charts - set(chart_paths))
            for chart_ref in missing_charts:
                warnings.append(f"Chart skipped because no valid plot points were rendered: {chart_ref}")
            if missing_charts:
                event_log.emit(
                    "charts.skipped",
                    level="warning",
                    chart_refs=missing_charts,
                )

        if request.template_id:
            try:
                ppt_path = output_dir / f"{job_id}.pptx"
                render_ppt(
                    bundle,
                    ppt_sections or sections,
                    request.template_id,
                    str(ppt_path),
                    images=chart_paths,
                    deck_plan=deck_plan,
                )
                artifacts["pptx"] = str(ppt_path)
                event_log.emit("artifact.written", artifact_type="pptx", path=str(ppt_path))
            except Exception as exc:
                logger.exception("PPT rendering failed")
                warnings.append(f"PPT rendering skipped: {exc}")
                event_log.emit(
                    "artifact.failed",
                    level="warning",
                    artifact_type="pptx",
                    error_type=type(exc).__name__,
                    error=str(exc),
                )

        if request.include_pdf:
            try:
                from app.renderers.pdf import render_pdf

                pdf_path = output_dir / f"{job_id}.pdf"
                render_pdf(
                    bundle=bundle,
                    sections=sections,
                    html_renderer=html_renderer,
                    language=request.language,
                    output_path=str(pdf_path),
                )
                artifacts["pdf"] = str(pdf_path)
                event_log.emit("artifact.written", artifact_type="pdf", path=str(pdf_path))
            except Exception as exc:
                logger.exception("PDF rendering failed")
                warnings.append(f"PDF rendering skipped: {exc}")
                event_log.emit(
                    "artifact.failed",
                    level="warning",
                    artifact_type="pdf",
                    error_type=type(exc).__name__,
                    error=str(exc),
                )

        event_log.emit(
            "job.completed",
            status=status,
            artifact_types=sorted(artifacts.keys()),
            warning_count=len(warnings),
        )
        return ReportJobResponse(
            job_id=job_id,
            batch_id=batch_id,
            status=status,
            bundle=bundle.model_dump(),
            sections=sections,
            ppt_sections=ppt_sections,
            llm_metadata=llm_metadata,
            artifacts=artifacts,
            warnings=warnings,
            logs_url=f"/api/v1/jobs/{job_id}/logs",
        )
    except Exception as exc:
        event_log.emit(
            "job.exception",
            level="error",
            error_type=type(exc).__name__,
            error=str(exc),
        )
        raise


@app.post("/api/v1/jobs", response_model=ReportJobResponse)
async def create_report_job(request: ReportJobRequest):
    """
    Demo-ready synchronous job path.

    The production design can move this behind the Postgres queue, but for a POC
    recording it proves the full Bundle -> Narrative -> HTML/PPT flow in one call.
    """
    return await execute_report_job(request)
    
@app.get("/api/v1/jobs/demo", response_model=ReportJobResponse)
async def demo_job():
    req = ReportJobRequest(
        wave_id="DEMO_WAVE_001",
        language="en-US",
        template_id="demo_master"
    )
    return await execute_report_job(req)
    
@app.get("/api/v1/wave-data")
def get_wave_data():
    from app.data.provider import get_provider
    provider = get_provider(settings)
    wave_id = provider.get_latest_wave_id()
    return provider.get_wave_data(wave_id)


@app.post("/api/v1/batches", response_model=BatchReportJobResponse)
async def create_report_batch(request: BatchReportJobRequest):
    ensure_template_registry_loaded()
    languages = dedupe_preserving_order(request.languages)
    if not languages:
        languages = ["en-US"]

    batch_id = str(uuid4())
    batch_log = StructuredEventLogger(
        settings.LOGS_DIR,
        batch_id=batch_id,
        context={
            "wave_id": request.wave_id,
            "template_id": request.template_id,
            "deck_mode": request.deck_mode,
        },
    )
    batch_log.emit(
        "batch.started",
        languages=languages,
        deck_mode=request.deck_mode,
        use_llm=request.use_llm,
        require_llm=request.require_llm,
        include_pdf=request.include_pdf,
    )

    child_job_ids = [str(uuid4()) for _ in languages]
    child_requests = [
        ReportJobRequest(
            wave_id=request.wave_id,
            language=language,
            template_id=request.template_id,
            deck_mode=request.deck_mode,
            allow_partial=request.allow_partial,
            audience=request.audience,
            tone=request.tone,
            use_llm=request.use_llm,
            require_llm=request.require_llm,
            include_pdf=request.include_pdf,
            output_dir=request.output_dir,
        )
        for language in languages
    ]
    results = await asyncio.gather(
        *[
            execute_report_job(child_request, batch_id=batch_id, job_id=job_id)
            for child_request, job_id in zip(child_requests, child_job_ids)
        ],
        return_exceptions=True,
    )

    children = []
    for language, job_id, result in zip(languages, child_job_ids, results):
        if isinstance(result, Exception):
            child_log = batch_log.child(job_id=job_id, context={"language": language})
            child_log.emit(
                "job.batch_exception",
                level="error",
                error_type=type(result).__name__,
                error=str(result),
            )
            children.append(
                BatchChildJobResponse(
                    language=language,
                    job_id=job_id,
                    status="failed",
                    warnings=[str(result)],
                    logs_url=f"/api/v1/jobs/{job_id}/logs",
                )
            )
            continue

        children.append(
            BatchChildJobResponse(
                language=language,
                job_id=result.job_id,
                status=result.status,
                artifacts=result.artifacts,
                warnings=result.warnings,
                llm_metadata=result.llm_metadata,
                logs_url=result.logs_url,
            )
        )

    status = "completed" if all(child.status == "completed" for child in children) else "failed"
    batch_log.emit(
        "batch.completed",
        status=status,
        child_statuses={child.language: child.status for child in children},
    )
    return BatchReportJobResponse(
        batch_id=batch_id,
        status=status,
        wave_id=request.wave_id,
        template_id=request.template_id,
        languages=languages,
        children=children,
        logs_url=f"/api/v1/batches/{batch_id}/logs",
    )


@app.get("/api/v1/batches/{batch_id}/logs")
def get_batch_logs(batch_id: str):
    return {"batch_id": batch_id, "events": read_jsonl(log_path("batches", batch_id))}


@app.get("/api/v1/jobs/{job_id}/logs")
def get_job_logs(job_id: str):
    return {"job_id": job_id, "events": read_jsonl(log_path("jobs", job_id))}
